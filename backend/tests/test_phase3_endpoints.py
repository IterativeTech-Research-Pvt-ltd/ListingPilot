import pytest
import io
import openpyxl
from unittest.mock import patch, AsyncMock
from fastapi.testclient import TestClient
from app.main import app
from app.models.product import Product
from app.models.draft import ListingDraft
from app.services.flipkart_adapter import FlipkartAdapter
from app.services.meesho_adapter import MeeshoAdapter
from app.services.rule_engine import RuleEngine

client = TestClient(app)

def test_flipkart_adapter():
    # 1. Credentials check (FK-1)
    status = FlipkartAdapter.validate_credentials()
    assert status["marketplace"] == "flipkart"
    assert "configured" in status
    assert status["mode"] in ("stub", "live")

    # 2. Stub listing creation (FK-1)
    draft_data = {
        "title": "Gulmohar Home Cotton Bedsheet",
        "bullets": ["100% Cotton", "Double Bedsheet"],
        "description": "Artisan bedsheet",
        "keywords": ["bedsheet", "cotton"]
    }
    product_data = {
        "sku": "JBP-101",
        "price": 1299.0,
        "stock": 40
    }
    import asyncio
    res = asyncio.run(FlipkartAdapter.create_listing(draft_data, product_data))
    assert res["success"] is True
    assert res["marketplace"] == "flipkart"
    assert res["external_listing_id"].startswith("LST")
    assert res["status"] == "live"

    # 3. Price & Stock update (FK-2)
    sync_res = asyncio.run(FlipkartAdapter.update_price_stock("JBP-101", 1399.0, 35))
    assert sync_res["success"] is True
    assert sync_res["updated_fields"]["price"] == 1399.0

def test_meesho_adapter():
    # 1. Status check
    status = MeeshoAdapter.validate_credentials()
    assert status["marketplace"] == "meesho"
    assert status["configured"] is True
    assert status["mode"] == "supplier_panel_files"

    # 2. Bulk template generation (MSH-1)
    sample_items = [
        {
            "sku": "JBP-101",
            "name": "Hand block print double bedsheet",
            "price": 1299.0,
            "mrp": 1699.0,
            "attributes": {"Material": "100% Cotton", "Size": "90x100 in"},
            "images": ["https://example.com/img1.jpg"],
            "draft": {
                "title": "Gulmohar Home Cotton Bedsheet",
                "bullets": ["Soft cotton", "Double bed"],
                "description": "Hand block print bedsheet"
            }
        }
    ]
    buf = MeeshoAdapter.generate_bulk_upload_file("Bedsheets", sample_items)
    assert isinstance(buf, io.BytesIO)
    
    # Verify openpyxl can load the generated workbook
    wb = openpyxl.load_workbook(buf)
    assert "Validate Sheet" in wb.sheetnames
    assert "Instructions" in wb.sheetnames
    assert "Example Sheet" in wb.sheetnames
    ws = wb["Validate Sheet"]
    header_values = [cell.value for cell in ws[1]]
    assert "Product Name *" in header_values
    assert "Seller SKU ID *" in header_values
    assert "Meesho Price (₹) *" in header_values
    assert "Wrong/Defective Price (₹) *" in header_values

def test_multichannel_rule_engine():
    # Test title limits per marketplace:
    # Amazon <= 200, Flipkart <= 120, Meesho <= 100
    mock_product = Product(
        sku="TEST-SKU",
        name="Artisan Bed Cover",
        category="Bedsheets",
        brand="Gulmohar Home",
        price=1200.0,
        stock=10,
        image_count=3,
        attributes={"Material": "Cotton", "Size": "Double", "Thread count": "180 TC", "Care": "Machine wash"}
    )

    long_title = "A" * 110 # Valid for Amazon, Valid for Flipkart, EXCEEDS Meesho 100

    draft_meesho = {"title": long_title, "bullets": ["b1", "b2", "b3"], "description": "clean"}
    res_msh = RuleEngine.evaluate(mock_product, draft_meesho, marketplace="meesho")
    assert res_msh["checks_summary"]["title_length_ok"] is False
    assert any("Meesho" in i["message"] for i in res_msh["issues"])

    draft_amz = {"title": long_title, "bullets": ["b1", "b2", "b3", "b4", "b5"], "description": "clean"}
    res_amz = RuleEngine.evaluate(mock_product, draft_amz, marketplace="amazon")
    assert res_amz["checks_summary"]["title_length_ok"] is True

    # Test Meesho plain text check (No HTML)
    html_draft = {"title": "Short Title", "bullets": ["b1", "b2", "b3"], "description": "<p>Formatted <b>text</b></p>"}
    res_html = RuleEngine.evaluate(mock_product, html_draft, marketplace="meesho")
    assert res_html["has_errors"] is True
    assert any("HTML" in i["message"] for i in res_html["issues"])

def test_credentials_status_endpoint():
    res = client.get("/api/v1/drafts/credentials-status")
    assert res.status_code == 200
    data = res.json()
    assert "amazon" in data
    assert "flipkart" in data
    assert "meesho" in data
    assert data["meesho"]["configured"] is True

def test_flipkart_publish_flow():
    # 1. Fetch a product
    prod_res = client.get("/api/v1/catalogue/products?page_size=1")
    assert prod_res.status_code == 200
    items = prod_res.json().get("items", [])
    if not items:
        pytest.skip("No products in database")
    product_id = items[0]["id"]

    # 2. Generate Flipkart Draft with mock
    with patch("app.services.ai_service.AIService.generate_listing", new_callable=AsyncMock) as mock_gen:
        mock_gen.return_value = {
            "title": "Gulmohar Home Cotton Bedsheet (Indigo Floral)",
            "bullets": ["100% Cotton", "180 TC", "Dimensions 90x100 in", "Pack of 1"],
            "description": "Artisanal hand block print bedsheet.",
            "keywords": ["flipkart bedsheet", "gulmohar home"]
        }
        gen_res = client.post(f"/api/v1/drafts/generate/{product_id}?marketplace=flipkart")
        assert gen_res.status_code == 200
        draft = gen_res.json()
        assert draft["marketplace"] == "flipkart"

    draft_id = draft["id"]

    # 3. Approve draft
    appr_res = client.post(f"/api/v1/drafts/{draft_id}/approve")
    assert appr_res.status_code == 200
    assert appr_res.json()["status"] == "approved"

    # 4. Publish to Flipkart (FK-1)
    pub_res = client.post(f"/api/v1/drafts/{draft_id}/publish")
    assert pub_res.status_code == 200
    pub_data = pub_res.json()
    assert pub_data["draft"]["status"] == "live"
    assert pub_data["draft"]["external_listing_id"].startswith("LST")

def test_meesho_bulk_export_endpoint():
    res = client.get("/api/v1/drafts/meesho/export")
    assert res.status_code == 200
    assert res.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    
    # Verify openpyxl can load the streamed content
    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    assert "Validate Sheet" in wb.sheetnames
    assert "Instructions" in wb.sheetnames
    assert "Example Sheet" in wb.sheetnames
    ws = wb["Validate Sheet"]
    header_values = [cell.value for cell in ws[1]]
    assert "Product Name *" in header_values
    assert "Meesho Price (₹) *" in header_values


def test_meesho_templates_endpoint():
    res = client.get("/api/v1/drafts/meesho/templates")
    assert res.status_code == 200
    data = res.json()
    assert data["success"] is True
    assert data["count"] >= 4
    categories = [t["category"] for t in data["templates"]]
    assert "Bedsheets" in categories
    assert "Shirts" in categories
    assert "Sneakers" in categories
    assert "Tshirts" in categories


def test_meesho_categories_endpoint():
    res = client.get("/api/v1/drafts/meesho/categories?search=Bedsheets")
    assert res.status_code == 200
    data = res.json()
    assert data["success"] is True
    assert data["count"] > 0
    matched = [c for c in data["categories"] if c["template_available"]]
    assert len(matched) > 0
    assert "Bedsheets" in matched[0]["name"]


def test_meesho_zip_export_endpoint():
    import zipfile
    res = client.get("/api/v1/drafts/meesho/export?as_zip=true")
    assert res.status_code == 200
    assert res.headers["content-type"] == "application/zip"
    zf = zipfile.ZipFile(io.BytesIO(res.content))
    namelist = zf.namelist()
    assert len(namelist) > 0
    assert any(n.endswith(".xlsx") for n in namelist)
