"""Meesho export: values come only from product data, gaps block the export."""
import asyncio
import io
import zipfile
from unittest.mock import AsyncMock, patch

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app.core.config import settings
from app.main import app
from app.models.product import Product
from app.services.ai_service import AIService
from app.services.catalogue_parser import CatalogueParser
from app.services.meesho_adapter import MeeshoAdapter, MeeshoExportIncomplete
from app.services.rule_engine import RuleEngine
from app.core.database import get_db

client = TestClient(app)

LEGAL = {
    "Country of Origin": "India",
    "Manufacturer Name": "Rangrez Textiles",
    "Manufacturer Address": "12 Udyog Vihar, Tiruppur, Tamil Nadu",
    "Manufacturer Pincode": "641604",
    "Packer Name": "Rangrez Textiles",
    "Packer Address": "12 Udyog Vihar, Tiruppur, Tamil Nadu",
    "Packer Pincode": "641604",
    "Importer Name": "Rangrez Textiles",
    "Importer Address": "12 Udyog Vihar, Tiruppur, Tamil Nadu",
    "Importer Pincode": "641604",
}


def tshirt(sku, title, **attr_overrides):
    attrs = {
        **LEGAL,
        "Size": "M",
        "GST %": "5",
        "HSN ID": "61091000",
        "Net Weight (gms)": "180",
        "Color": "Black",
        "Material": "100% cotton",
        "Fit/Shape": "Regular",
        "Neck": "Round",
        "Net Quantity (N)": "1",
        "Occasion": "Casual",
        "Pattern": "Solid",
        "Print Or Pattern Type": "Solid",
        "Sleeve Length": "Short Sleeves",
        "Chest Size": "40",
        "Length Size": "27",
        "Shoulder Size": "17",
    }
    attrs.update(attr_overrides)
    attrs = {k: v for k, v in attrs.items() if v is not None}
    return {
        "sku": sku, "name": title, "category": "Tshirts", "brand": "Gulmohar Home",
        "price": 499.0, "mrp": 799.0, "stock": 12,
        "attributes": attrs,
        "images": ["https://example.com/front.jpg"],
        "draft": {"title": title, "bullets": ["SOFT COTTON — breathable"], "description": "Plain black tee."},
    }


def sheet_row(buf, row=5):
    tmpl = MeeshoAdapter.find_template("Tshirts")
    ws = openpyxl.load_workbook(buf)[tmpl["sheet_name"]]
    return {name: ws.cell(row=row, column=col).value for col, name in tmpl["columns"].items()}


def export_problems(items, category="Tshirts"):
    with pytest.raises(MeeshoExportIncomplete) as exc:
        MeeshoAdapter.generate_batch_export(items, category_filter=category)
    return exc.value.problems


def test_complete_product_exports_its_own_values():
    buf, _, _ = MeeshoAdapter.generate_batch_export([tshirt("TS-1", "Gulmohar Home Black Cotton Tee")], category_filter="Tshirts")
    row = sheet_row(buf)
    assert row["Product Name"] == "Gulmohar Home Black Cotton Tee"
    assert row["Meesho Price"] == 499.0  # the seller's price, not a computed discount
    assert row["MRP"] == 799.0
    assert row["HSN ID"] == "61091000"
    assert row["Fabric"] == "Cotton"  # "100% cotton" names exactly one option
    assert row["Manufacturer Name"] == "Rangrez Textiles"
    assert row["Net Weight (gms)"] == 180
    assert row["Wrong/Defective Return Discount(₹)"] is None  # optional and not in the data
    assert row["Brand Name"] == "Gulmohar Home"
    assert row["Brand"] is None  # brand dropdown is never auto-picked


def test_missing_compulsory_fields_block_export_instead_of_defaults():
    problems = export_problems([tshirt("TS-2", "Tee two", **{"HSN ID": None, "Manufacturer Name": None, "Net Weight (gms)": None})])
    fields = {p["field"] for p in problems if p["blocking"]}
    assert {"HSN ID", "Manufacturer Name", "Net Weight (gms)"} <= fields


def test_values_outside_the_dropdown_are_not_snapped():
    # Old behaviour: HSN 6109 fell through to the first option and unknown sizes became a default
    problems = export_problems([tshirt("TS-3", "Tee three", **{"HSN ID": "6109", "Size": "Universal"})])
    by_field = {p["field"]: p for p in problems}
    assert by_field["HSN ID"]["blocking"] and "'6109'" in by_field["HSN ID"]["reason"]
    assert by_field["Variation"]["blocking"]


def test_shared_product_name_across_designs_blocks_export():
    items = [tshirt("TS-4", "Same tee title"), tshirt("TS-5", "Same tee title", Color="White")]
    problems = export_problems(items)
    assert {p["sku"] for p in problems if p["field"] == "Product Name"} == {"TS-4", "TS-5"}


def test_size_variants_of_one_style_may_share_a_name():
    items = [tshirt("TS-6M", "Variant tee", **{"Style ID": "TS-6"}),
             tshirt("TS-6L", "Variant tee", **{"Style ID": "TS-6", "Size": "L"})]
    buf, _, _ = MeeshoAdapter.generate_batch_export(items, category_filter="Tshirts")
    assert sheet_row(buf, 5)["Variation"] == "M" and sheet_row(buf, 6)["Variation"] == "L"


@pytest.mark.parametrize("val, options, col, expected", [
    ("180 TC", ["100", "180", "200"], "Thread Count", "180"),
    ("5.0", ["18", "5"], "GST %", "5"),
    ("100% cotton", ["Cotton", "Cotton Blend", "Polyester"], "Fabric", "Cotton"),
    ("Cotton blend fabric", ["Cotton", "Cotton Blend"], "Fabric", "Cotton Blend"),
    ("90 x 100 in (double)", ["Double", "King", "Single"], "Variation", "Double"),
    ("Double / King, 90 x 108 in", ["Double", "King"], "Variation", None),  # ambiguous
    ("Floral, hand block print", ["Batik", "Floral", "Geometric"], "Print or Pattern Type", "Floral"),
    ("Indigo and white", ["Blue", "White"], "Color", None),  # colour must match exactly
    ("6109", ["392410", "6304"], "HSN ID", None),
    ("Gulmohar", ["Gulmohar Home Decor"], "Brand", None),
    ("group  1", ["Group  1", "Group  10"], "Group ID", "Group  1"),  # template spelling kept
])
def test_dropdown_matching_is_high_confidence_only(val, options, col, expected):
    assert MeeshoAdapter._match_dropdown_option(val, options, col_name=col) == expected


def bedsheet(sku, title, **attr_overrides):
    attrs = {
        **LEGAL,
        "Material": "100% cotton",
        "Size": "90 x 100 in (double)",
        "Thread count": "180 TC",
        "Colour": "Indigo and white",
        "Color": "Blue",
        "Pattern": "Floral, hand block print",
        "Included": "1 bedsheet + 2 pillow covers (17 x 27 in)",
        "GST %": "5",
        "HSN ID": "63041910",
        "Net Weight (gms)": "650",
        "Net Quantity (N)": "1",
        # Facts the seller states for Meesho; none can be read off the measurements above
        "Bed Size": "Double Queen",
        "Care Instructions": "Machine Wash Only",
        "Ideal For": "Adult",
        "Ideal Season": "Summer",
        "Occassion Type": "Others",
        "Quality": "Regular",
        "Type": "Flat Sheets",
    }
    attrs.update(attr_overrides)
    item = tshirt(sku, title)
    item.update(category="Bedsheets", attributes={k: v for k, v in attrs.items() if v is not None})
    return item


def bedsheet_row(buf, row=5):
    tmpl = MeeshoAdapter.find_template("Bedsheets")
    ws = openpyxl.load_workbook(buf)[tmpl["sheet_name"]]
    return {name: ws.cell(row=row, column=col).value for col, name in tmpl["columns"].items()}


def test_bedsheet_measurements_come_from_the_stated_sizes():
    buf, _, _ = MeeshoAdapter.generate_batch_export([bedsheet("BS-1", "Indigo floral bedsheet")], category_filter="Bedsheets")
    row = bedsheet_row(buf)
    assert (row["Variation"], row["Thread Count"]) == ("Double", "180")
    assert (row["Product Length"], row["Product Width"], row["Product Dimension Unit"]) == (100, 90, "inch")
    assert row["Dimensions (inches)"] == "90*100-(Double)"  # sides and Variation pick one option
    assert row["No. Of Pillow Covers"] == "2"
    assert row["Size"] == "Double Queen"  # from Bed Size; Size itself holds the measurement
    assert row["Color"] == "Blue"  # the seller's Meesho colour wins over the descriptive Colour
    assert row["Print or Pattern Type"] == "Floral"


def test_bed_size_is_never_guessed_from_the_measurement():
    # "(double)" could be Double Queen or Double King; the export blocks and names the attribute to set
    problems = export_problems([bedsheet("BS-6", "Unsized bedsheet", **{"Bed Size": None})], category="Bedsheets")
    by_field = {p["field"]: p for p in problems if p["blocking"]}
    assert set(by_field) == {"Size"}
    buf = MeeshoAdapter.build_fill_in_sheet([bedsheet("BS-6", "Unsized bedsheet", **{"Bed Size": None})])
    headers = [c.value for c in openpyxl.load_workbook(buf)["Fill in"][1]]
    assert "Bed Size" in headers and "Size" not in headers  # importing it can't overwrite the measurement


@pytest.mark.parametrize("size, variation, expected", [
    ("90 x 100 in (double)", "Double", "90*100-(Double)"),
    ("100 x 90 in", "Queen", "100*90-(Queen)"),
    ("90 x 100 in", None, None),  # two options share these sides; no Variation to choose between them
    ("60 x 61 in", "Single", None),  # no option with these sides
    ("Double bed", "Double", None),  # no stated sides
])
def test_dimension_option_needs_the_stated_sides(size, variation, expected):
    from app.services.meesho_adapter import _dimension_option
    options = ["100*90-(Queen)", "90*100-(Double)", "108*108-(King)"]
    assert _dimension_option(size, variation, options) == expected


def test_official_workbook_gets_no_extra_sheets():
    buf, _, _ = MeeshoAdapter.generate_batch_export([bedsheet("BS-2", "Plain bedsheet")], category_filter="Bedsheets")
    master = openpyxl.load_workbook(MeeshoAdapter.find_template("Bedsheets")["filepath"])
    assert openpyxl.load_workbook(buf).sheetnames == master.sheetnames


def test_category_without_official_template_is_refused():
    item = tshirt("PL-1", "Ceramic planter")
    item["category"] = "Planters"
    problems = export_problems([item], category="Planters")
    assert [(p["sku"], p["field"]) for p in problems] == [("PL-1", "Template")]


def test_ethnic_footwear_is_not_sent_to_the_sneakers_template():
    assert MeeshoAdapter.find_template("Ethnic footwear") is None


@pytest.mark.parametrize("category, expected", [
    ("Double Bedsheets", "Bedsheets"),
    ("Bed Sheets", "Bedsheets"),
    ("T-Shirts", "Tshirts"),
    ("Tees", "Tshirts"),
    ("Brother & Sister T-Shirts", "Tshirts"),
    ("Casual Shirts", "Shirts"),
    ("Men Fashion/Footwear/Shoes/Sneakers", "Sneakers"),
    # Words that merely contain a template's name match nothing
    ("Steel bottles", None),
    ("Diwans & Settees", None),
    ("Teeth Whitening", None),
    ("Sweatshirts", None),
    ("Shirt Fabric", None),
    ("Men", None),
    ("Pet Bedding", None),
])
def test_categories_match_templates_by_whole_words(category, expected):
    tmpl = MeeshoAdapter.find_template(category)
    assert (tmpl["category"] if tmpl else None) == expected


def test_sole_material_and_weave_pattern_are_not_filled_from_other_attributes():
    tmpl = MeeshoAdapter.find_template("Sneakers")
    row = {"attrs": {"Material": "Canvas", "Pattern": "Solid"}}
    assert MeeshoAdapter._template_value("material", "Material", row, "Sneakers") == "Canvas"
    assert MeeshoAdapter._template_value("sole material", "Sole Material", row, "Sneakers") is None
    assert MeeshoAdapter._template_value("weave pattern", "Weave Pattern", row, "Shirts") is None
    assert MeeshoAdapter._template_value("pattern", "Pattern", row, "Shirts") == "Solid"
    assert tmpl is not None


def test_several_categories_always_zip():
    buf, media_type, _ = MeeshoAdapter.generate_batch_export([tshirt("TS-7", "Zip tee"), bedsheet("BS-3", "Zip bedsheet")])
    assert media_type == "application/zip"
    assert sorted(zipfile.ZipFile(buf).namelist()) == ["ListingPilot_Meesho_Bedsheets.xlsx", "ListingPilot_Meesho_Tshirts.xlsx"]


def test_fill_in_sheet_round_trips_through_catalogue_import():
    item = bedsheet("BS-4", "Round trip bedsheet", **{"HSN ID": None, "Color": None, "Manufacturer Pincode": None})
    wb = openpyxl.load_workbook(MeeshoAdapter.build_fill_in_sheet([item]))
    ws = wb["Fill in"]
    headers = [c.value for c in ws[1]]
    assert {"HSN ID", "Color", "Manufacturer Pincode"} <= set(headers)
    assert ws.data_validations.dataValidation  # dropdowns carry the template's own options

    # Seller fills the gaps, the sheet is imported, and the export now passes
    filled = {"HSN ID": "63041910", "Color": "Blue", "Manufacturer Pincode": "641604"}
    for name, value in filled.items():
        ws.cell(row=2, column=headers.index(name) + 1, value=value)
    out = io.BytesIO()
    wb.save(out)
    parsed, errors, _ = CatalogueParser.parse_file("fill.xlsx", out.getvalue())
    assert not errors and parsed[0]["sku"] == "BS-4"
    item["attributes"].update(parsed[0]["attributes"])
    item["images"] = parsed[0]["images"]
    MeeshoAdapter.generate_batch_export([item], category_filter="Bedsheets")


def test_fill_in_sheet_is_none_when_nothing_is_missing():
    assert MeeshoAdapter.build_fill_in_sheet([bedsheet("BS-5", "Complete bedsheet")]) is None


def test_export_endpoint_reports_problems_with_422():
    res = client.get("/api/v1/drafts/meesho/export?category=Bedsheets")
    assert res.status_code == 422
    detail = res.json()["detail"]
    assert "blocked" in detail["message"]
    assert any(p["sku"] == "JBP-101" for p in detail["problems"])


def test_export_endpoint_streams_complete_category():
    item = tshirt("TS-API-1", "Endpoint tee")
    db_gen = app.dependency_overrides[get_db]()  # the test database set up in conftest
    db = next(db_gen)
    try:
        if not db.query(Product).filter(Product.sku == "TS-API-1").first():
            db.add(Product(**{k: item[k] for k in ["sku", "name", "category", "brand", "price", "mrp", "stock", "attributes", "images"]}))
            db.commit()
    finally:
        db_gen.close()
    res = client.get("/api/v1/drafts/meesho/export?category=Tshirts")
    assert res.status_code == 200, res.text
    assert sheet_row(io.BytesIO(res.content))["SKU ID"] == "TS-API-1"


def test_rule_engine_warns_on_words_meesho_flags():
    product = Product(sku="X", name="x", category="Bedsheets", price=100, attributes={}, images=["a"], image_count=1)
    draft = {"title": "A cotton bedsheet set", "bullets": ["CARE — machine wash"], "description": "Soft comfort."}
    issues = RuleEngine.evaluate(product, draft, marketplace="meesho")["issues"]
    flagged = [i for i in issues if i["kind"] == "meesho_flagged_word"]
    assert flagged and flagged[0]["level"] == "warning"
    assert not [i for i in RuleEngine.evaluate(product, draft, marketplace="amazon")["issues"] if i["kind"] == "meesho_flagged_word"]


def test_required_attributes_match_whatever_case_the_import_stored():
    # The importer title-cases headers: a sheet's "Thread count" column is stored as "Thread Count"
    attrs = {"Material": "Cotton", "Size": "90 x 100 in", "Thread Count": "180 TC", "Care": "Machine wash"}
    product = Product(sku="X", name="x", category="Bedsheets", price=100, attributes=attrs, images=["a"], image_count=1)
    issues = RuleEngine.evaluate(product, None, marketplace="meesho")["issues"]
    assert not [i for i in issues if i["kind"] == "attribute"]


def test_meesho_prompt_and_template_copy_avoid_flagged_words(monkeypatch):
    product = Product(sku="JBP-1", name="Bedsheet", category="Bedsheets", brand="Gulmohar Home", price=1299,
                      attributes={"Colour": "Blue", "Care": "Machine wash cold"}, images=[])
    monkeypatch.setattr(settings, "AI_PROVIDER", "groq")
    monkeypatch.setattr(settings, "GROQ_API_KEY", "gsk_test_key")
    fake = AsyncMock(return_value={"title": "t"})
    with patch.object(AIService, "_call_chat_completions", fake):
        asyncio.run(AIService.generate_listing(product, marketplace="meesho"))
        asyncio.run(AIService.generate_listing(product, marketplace="amazon"))
    meesho_prompt, amazon_prompt = fake.call_args_list[0].args[5], fake.call_args_list[1].args[5]
    assert "Meesho rules" in meesho_prompt and "WASH & MAINTENANCE" in meesho_prompt
    assert "Meesho rules" not in amazon_prompt

    bullets = AIService._fallback_generator(product, "Gulmohar Home", marketplace="meesho")["bullets"]
    assert not any(b.upper().startswith("CARE") for b in bullets)
