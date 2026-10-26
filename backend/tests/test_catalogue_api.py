import pytest
from fastapi.testclient import TestClient
import io
import uuid
import openpyxl
from app.main import app

client = TestClient(app)

def test_health_check():
    response = client.get("/api/v1/health")
    assert response.status_code == 200
    data = response.json()
    assert data["status"] == "healthy"
    assert data["phase"] == 1

def test_download_template():
    response = client.get("/api/v1/catalogue/template")
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    
    # Verify openpyxl can load the generated workbook
    wb = openpyxl.load_workbook(io.BytesIO(response.content))
    assert "Master Catalogue" in wb.sheetnames
    assert "Sample Pilot Products" in wb.sheetnames
    assert "Instructions & Rules" in wb.sheetnames
    
    ws = wb["Master Catalogue"]
    assert ws.cell(row=1, column=1).value == "SKU *"
    assert ws.cell(row=1, column=2).value == "Product Name *"

def test_list_products():
    response = client.get("/api/v1/catalogue/products")
    assert response.status_code == 200
    data = response.json()
    assert "items" in data
    assert "total" in data
    assert data["total"] >= 50
    assert len(data["categories"]) >= 5

def test_stats_endpoint():
    response = client.get("/api/v1/catalogue/stats")
    assert response.status_code == 200
    data = response.json()
    assert data["total_skus"] >= 50
    assert data["total_stock"] > 0
    assert data["avg_price"] > 0
    assert "category_counts" in data

def test_product_crud():
    unique_sku = f"TEST-SKU-{uuid.uuid4().hex[:6].upper()}"
    new_prod = {
        "sku": unique_sku,
        "name": "Artisan Handcrafted Brass Bell",
        "category": "Pooja items",
        "brand": "Gulmohar Home",
        "mrp": 899.0,
        "price": 649.0,
        "stock": 15,
        "images": ["https://images.unsplash.com/test.jpg"],
        "attributes": {
            "Material": "Solid Brass",
            "Weight": "250 g"
        },
        "description": "Test description for bell"
    }
    create_res = client.post("/api/v1/catalogue/products", json=new_prod)
    assert create_res.status_code == 201, create_res.text
    created_id = create_res.json()["id"]

    # 2. Get product
    get_res = client.get(f"/api/v1/catalogue/products/{created_id}")
    assert get_res.status_code == 200
    assert get_res.json()["sku"] == unique_sku

    # 3. Update product
    update_res = client.put(f"/api/v1/catalogue/products/{created_id}", json={
        "stock": 25,
        "price": 599.0
    })
    assert update_res.status_code == 200
    assert update_res.json()["stock"] == 25
    assert update_res.json()["price"] == 599.0

    # 4. Delete product
    del_res = client.delete(f"/api/v1/catalogue/products/{created_id}")
    assert del_res.status_code == 200

    # 5. Verify deleted
    get_after_del = client.get(f"/api/v1/catalogue/products/{created_id}")
    assert get_after_del.status_code == 404
