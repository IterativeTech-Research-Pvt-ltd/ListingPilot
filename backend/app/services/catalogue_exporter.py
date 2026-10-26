import io
import csv
import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from typing import List
from app.models.product import Product

class CatalogueExporter:
    """Exports master catalogue products to Excel, CSV, or JSON for client data portability."""

    @classmethod
    def export_excel(cls, products: List[Product]) -> io.BytesIO:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Master Catalogue Export"
        ws.views.sheetView[0].showGridLines = True

        # Styles
        header_fill = PatternFill(start_color="0D6B58", end_color="0D6B58", fill_type="solid")
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        regular_font = Font(name="Segoe UI", size=10)
        thin_border = Border(
            left=Side(style='thin', color='DBE1DC'),
            right=Side(style='thin', color='DBE1DC'),
            top=Side(style='thin', color='DBE1DC'),
            bottom=Side(style='thin', color='DBE1DC')
        )

        # Collect all dynamic attribute keys across all products
        attr_keys_set = set()
        for p in products:
            if p.attributes and isinstance(p.attributes, dict):
                attr_keys_set.update(p.attributes.keys())
        sorted_attr_keys = sorted(list(attr_keys_set))

        headers = [
            "SKU", "Product Name", "Category", "Brand", "MRP (INR)", "Selling Price (INR)",
            "Stock", "Image Count", "Image URLs"
        ] + sorted_attr_keys + ["Description", "Created At"]

        # Write header row
        for col_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=h)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center" if col_idx in [1, 5, 6, 7, 8] else "left", vertical="center")
            cell.border = thin_border

        ws.row_dimensions[1].height = 26

        # Write data rows
        for row_idx, p in enumerate(products, start=2):
            image_urls_str = ", ".join(p.images) if p.images else ""
            img_count = p.image_count or (len(p.images) if p.images else 0)

            row_data = [
                p.sku,
                p.name,
                p.category,
                p.brand,
                p.mrp,
                p.price,
                p.stock,
                img_count,
                image_urls_str
            ]

            # Dynamic attributes
            p_attrs = p.attributes or {}
            for k in sorted_attr_keys:
                row_data.append(p_attrs.get(k, ""))

            row_data.append(p.description or "")
            row_data.append(p.created_at.strftime("%Y-%m-%d %H:%M:%S") if p.created_at else "")

            for col_idx, val in enumerate(row_data, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.font = regular_font
                cell.border = thin_border
                if col_idx in [5, 6] and isinstance(val, (int, float)):
                    cell.number_format = '₹#,##0'
                    cell.alignment = Alignment(horizontal="right")
                elif col_idx in [1, 7, 8]:
                    cell.alignment = Alignment(horizontal="center")

        # Auto width
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or '')
                max_len = max(max_len, min(len(val_str), 40))
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    @classmethod
    def export_csv(cls, products: List[Product]) -> io.BytesIO:
        output = io.StringIO()
        writer = csv.writer(output)

        attr_keys_set = set()
        for p in products:
            if p.attributes and isinstance(p.attributes, dict):
                attr_keys_set.update(p.attributes.keys())
        sorted_attr_keys = sorted(list(attr_keys_set))

        headers = [
            "SKU", "Product Name", "Category", "Brand", "MRP", "Selling Price",
            "Stock", "Image Count", "Image URLs"
        ] + sorted_attr_keys + ["Description"]

        writer.writerow(headers)

        for p in products:
            p_attrs = p.attributes or {}
            img_urls = ", ".join(p.images) if p.images else ""
            img_count = p.image_count or (len(p.images) if p.images else 0)
            row = [
                p.sku,
                p.name,
                p.category,
                p.brand,
                p.mrp or "",
                p.price,
                p.stock,
                img_count,
                img_urls
            ]
            for k in sorted_attr_keys:
                row.append(p_attrs.get(k, ""))
            row.append(p.description or "")
            writer.writerow(row)

        output_bytes = io.BytesIO(output.getvalue().encode("utf-8-sig"))
        output_bytes.seek(0)
        return output_bytes

    @classmethod
    def export_json(cls, products: List[Product]) -> str:
        data = [p.to_dict() for p in products]
        return json.dumps(data, indent=2, ensure_ascii=False)
